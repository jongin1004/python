from openpyxl import load_workbook
wb = load_workbook('sample.xlsx')
ws = wb.active

from openpyxl.chart import BarChart, LineChart, Reference

# # 데이터 범위를 지정
# bar_value = Reference(ws, min_row = 2, max_row = 11, min_col = 2, max_col = 3)
# # 차트 종류 설정 ( Bar, Line, Pie, ...)
# bar_chart = BarChart() 
# # 차트에 위의 만든 데이터를 추가
# bar_chart.add_data(bar_value)

# # sheet 지정한 위치에, 위의 생성한 chart를 삽입
# ws.add_chart(bar_chart, "E1") 

# B1:C11 까지의 데이터 
line_value = Reference(ws, min_row = 1, max_row = 11, min_col = 2, max_col = 3)
line_chart = LineChart()
line_chart.add_data(line_value, titles_from_data = True) # 항목명을 가져와서 계열로 설정
line_chart.title = '성적표'
line_chart.style = 5 # 미리 정의된 스타일을 적용
line_chart.y_axis.title = '점수'
line_chart.x_axis.title = '번호'
ws.add_chart(line_chart, 'E1')

wb.save('sample_modify.xlsx')
wb.close()
